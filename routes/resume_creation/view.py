from __future__ import annotations

import io
from typing import Any, Dict, List, Optional

from flask import jsonify, render_template, request, session, send_file
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas

from routes.auth import login_required
from utils import build_session_context
from services.resume_creation.resume_service import (
    ALLOWED_TEMPLATES,
    build_content_from_analysis,
    compute_completeness,
    create_resume,
    create_resume_version,
    delete_resume,
    duplicate_resume,
    get_owned_resume,
    get_owned_resume_version,
    list_versions_for_resume,
    list_resumes_for_user,
    restore_resume_from_version,
    sanitize_content,
    update_resume,
)
from services.resume_intelligence.ats import analyze_ats
from services.resume_intelligence.generative.local_generation import (
    TASK_BULLETS,
    TASK_EXPERIENCE,
    TASK_RESUME_REWRITE,
    TASK_SUMMARY,
    local_generate,
)
from services.resume_intelligence.generative.validators import validate_generated_claims
from services.resume_intelligence.matching.job_matcher import match_resume_to_job


@login_required
def render_resume_creation():
    return render_template("resume_creation.html", **build_session_context(session))


def _api_error(status: int, message: str, *, code: str = "error"):
    return jsonify({"success": False, "error": message, "code": code}), status


def _require_api_user_id() -> Optional[int]:
    raw_user_id = session.get("user_id")
    if raw_user_id is None:
        return None
    try:
        return int(raw_user_id)
    except (TypeError, ValueError):
        return None


def _resume_payload(resume) -> Dict[str, Any]:
    payload = resume.to_dict()
    payload["completeness"] = compute_completeness(payload.get("content") or {})
    return payload


def _filename_safe(value: str) -> str:
    safe = "".join(char for char in str(value or "resume") if char.isalnum() or char in {"-", "_", " "}).strip()
    safe = "-".join(safe.split())
    return safe[:80] or "resume"


def _content_to_text(content: Dict[str, Any]) -> str:
    personal = content.get("personal_information") or {}
    lines: List[str] = []

    for key in ("full_name", "professional_title", "email", "phone", "location"):
        value = str(personal.get(key) or "").strip()
        if value:
            lines.append(value)

    summary = str(content.get("summary") or "").strip()
    if summary:
        lines.append("Summary")
        lines.append(summary)

    for section_title, section_key in (
        ("Experience", "experience"),
        ("Education", "education"),
        ("Projects", "projects"),
        ("Certifications", "certifications"),
    ):
        rows = content.get(section_key) or []
        if not rows:
            continue
        lines.append(section_title)
        for row in rows:
            if not isinstance(row, dict):
                continue
            for value in row.values():
                if isinstance(value, list):
                    for item in value:
                        item_text = str(item).strip()
                        if item_text:
                            lines.append(item_text)
                else:
                    text = str(value).strip()
                    if text:
                        lines.append(text)

    skills = content.get("skills") or []
    if skills:
        lines.append("Skills")
        lines.append(", ".join(str(skill).strip() for skill in skills if str(skill).strip()))

    languages = content.get("languages") or []
    if languages:
        lines.append("Languages")
        for language in languages:
            if isinstance(language, dict):
                label = str(language.get("language") or "").strip()
                proficiency = str(language.get("proficiency") or "").strip()
                if label:
                    lines.append(f"{label} {proficiency}".strip())

    custom_sections = content.get("custom_sections") or []
    for section in custom_sections:
        if not isinstance(section, dict):
            continue
        title = str(section.get("title") or "").strip()
        if title:
            lines.append(title)
        for item in section.get("items") or []:
            if not isinstance(item, dict):
                continue
            for value in item.values():
                if isinstance(value, list):
                    for bullet in value:
                        bullet_text = str(bullet).strip()
                        if bullet_text:
                            lines.append(bullet_text)
                else:
                    text = str(value).strip()
                    if text:
                        lines.append(text)

    return "\n".join(lines).strip()


def _content_to_sections(content: Dict[str, Any]) -> List[Dict[str, Any]]:
    sections: List[Dict[str, Any]] = []
    personal = content.get("personal_information") or {}

    header_lines = [
        str(personal.get("full_name") or "").strip(),
        str(personal.get("professional_title") or "").strip(),
        " | ".join(
            part for part in [
                str(personal.get("email") or "").strip(),
                str(personal.get("phone") or "").strip(),
                str(personal.get("location") or "").strip(),
            ] if part
        ),
    ]
    links = [
        str(personal.get("linkedin") or "").strip(),
        str(personal.get("github") or "").strip(),
        str(personal.get("portfolio") or "").strip(),
    ]
    link_line = " | ".join([item for item in links if item])
    if link_line:
        header_lines.append(link_line)

    sections.append({
        "title": "",
        "lines": [line for line in header_lines if line],
        "bullets": [],
    })

    summary = str(content.get("summary") or "").strip()
    if summary:
        sections.append({"title": "Summary", "lines": [summary], "bullets": []})

    def entry_to_lines(entry: Dict[str, Any], keys: List[str]) -> List[str]:
        lines: List[str] = []
        for key in keys:
            value = entry.get(key)
            if isinstance(value, list):
                continue
            text = str(value or "").strip()
            if text:
                lines.append(text)
        return lines

    def add_section(title: str, key: str, line_keys: List[str], bullets_key: Optional[str] = None):
        rows = content.get(key) or []
        if not rows:
            return
        lines: List[str] = []
        bullets: List[str] = []
        for row in rows:
            if not isinstance(row, dict):
                continue
            row_lines = entry_to_lines(row, line_keys)
            if row_lines:
                lines.append(" | ".join(row_lines))
            if bullets_key:
                for bullet in row.get(bullets_key) or []:
                    bullet_text = str(bullet or "").strip()
                    if bullet_text:
                        bullets.append(bullet_text)
        if lines or bullets:
            sections.append({"title": title, "lines": lines, "bullets": bullets})

    add_section("Experience", "experience", ["title", "company", "location", "start_date", "end_date"], bullets_key="bullets")
    add_section("Education", "education", ["degree", "field_of_study", "institution", "location", "start_date", "end_date", "grade"])

    skills = content.get("skills") or []
    if skills:
        sections.append({"title": "Skills", "lines": [", ".join(str(skill) for skill in skills if str(skill).strip())], "bullets": []})

    add_section("Projects", "projects", ["name", "role", "project_url", "github_url", "start_date", "end_date"], bullets_key=None)
    add_section("Certifications", "certifications", ["name", "issuer", "issue_date", "expiration_date", "credential_id", "credential_url"], bullets_key=None)
    add_section("Languages", "languages", ["language", "proficiency"], bullets_key=None)

    custom_sections = content.get("custom_sections") or []
    for custom in custom_sections:
        if not isinstance(custom, dict):
            continue
        title = str(custom.get("title") or "").strip()
        lines = []
        bullets = []
        for item in custom.get("items") or []:
            if not isinstance(item, dict):
                continue
            descriptor = " | ".join(
                token for token in [
                    str(item.get("title") or "").strip(),
                    str(item.get("subtitle") or "").strip(),
                    str(item.get("date") or "").strip(),
                ] if token
            )
            if descriptor:
                lines.append(descriptor)
            description = str(item.get("description") or "").strip()
            if description:
                lines.append(description)
            for bullet in item.get("bullets") or []:
                bullet_text = str(bullet).strip()
                if bullet_text:
                    bullets.append(bullet_text)
        if title and (lines or bullets):
            sections.append({"title": title, "lines": lines, "bullets": bullets})

    return sections


def _export_pdf_bytes(resume) -> bytes:
    content = sanitize_content(resume.get_content())
    sections = _content_to_sections(content)

    buffer = io.BytesIO()
    doc = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4

    y = height - 0.75 * inch
    left = 0.75 * inch
    line_height = 14

    def ensure_space(lines_needed: int = 1):
        nonlocal y
        if y - (lines_needed * line_height) < 0.75 * inch:
            doc.showPage()
            y = height - 0.75 * inch

    for idx, section in enumerate(sections):
        title = section.get("title") or ""
        lines = section.get("lines") or []
        bullets = section.get("bullets") or []

        if idx == 0:
            for i, line in enumerate(lines):
                if not line:
                    continue
                ensure_space()
                if i == 0:
                    doc.setFont("Helvetica-Bold", 15)
                elif i == 1:
                    doc.setFont("Helvetica", 11)
                else:
                    doc.setFont("Helvetica", 9.5)
                doc.drawString(left, y, line[:125])
                y -= line_height
            y -= 8
            continue

        if title:
            ensure_space(2)
            doc.setFont("Helvetica-Bold", 11)
            doc.drawString(left, y, title.upper()[:90])
            y -= line_height
            doc.setLineWidth(0.5)
            doc.line(left, y + 4, width - left, y + 4)

        doc.setFont("Helvetica", 9.7)
        for line in lines:
            if not line:
                continue
            ensure_space()
            doc.drawString(left, y, line[:130])
            y -= line_height

        for bullet in bullets:
            if not bullet:
                continue
            ensure_space()
            doc.drawString(left + 8, y, f"- {bullet}"[:128])
            y -= line_height

        y -= 6

    doc.save()
    buffer.seek(0)
    return buffer.getvalue()


def _export_docx_bytes(resume) -> bytes:
    content = sanitize_content(resume.get_content())
    sections = _content_to_sections(content)

    document = Document()

    for idx, section in enumerate(sections):
        title = section.get("title") or ""
        lines = section.get("lines") or []
        bullets = section.get("bullets") or []

        if idx == 0:
            if lines:
                heading = document.add_paragraph(lines[0])
                if heading.runs:
                    heading.runs[0].bold = True
                    heading.runs[0].font.size = None
            for line in lines[1:]:
                if line:
                    document.add_paragraph(line)
            continue

        if title:
            document.add_heading(title, level=2)
        for line in lines:
            if line:
                document.add_paragraph(line)
        for bullet in bullets:
            if bullet:
                document.add_paragraph(bullet, style="List Bullet")

    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output.getvalue()


def _get_resume_source_for_generation(content: Dict[str, Any], body: Dict[str, Any]) -> str:
    source_text = str(body.get("source_text") or "").strip()
    if source_text:
        return source_text

    source_type = str(body.get("source_type") or "").strip().lower()
    source_index = body.get("source_index")

    if source_type == "experience" and isinstance(source_index, int):
        rows = content.get("experience") or []
        if 0 <= source_index < len(rows) and isinstance(rows[source_index], dict):
            row = rows[source_index]
            bullets = row.get("bullets") or []
            return "\n".join(str(line).strip() for line in bullets if str(line).strip())

    if source_type == "project" and isinstance(source_index, int):
        rows = content.get("projects") or []
        if 0 <= source_index < len(rows) and isinstance(rows[source_index], dict):
            return str(rows[source_index].get("description") or "").strip()

    if source_type == "summary":
        return str(content.get("summary") or "").strip()

    return ""


def _apply_generated_content(resume, body: Dict[str, Any], generated_text: str):
    apply_to = str(body.get("apply_to") or "").strip().lower()
    source_index = body.get("source_index")

    content = resume.get_content()

    if apply_to == "summary":
        content["summary"] = generated_text.strip()
        update_resume(resume, content=content)
        return

    if apply_to == "experience" and isinstance(source_index, int):
        experience = content.get("experience") or []
        if 0 <= source_index < len(experience) and isinstance(experience[source_index], dict):
            bullet_lines = [
                line.strip("•- \t") for line in generated_text.splitlines() if line.strip("•- \t")
            ]
            experience[source_index]["bullets"] = bullet_lines
            content["experience"] = experience
            update_resume(resume, content=content)
            return

    if apply_to == "project" and isinstance(source_index, int):
        projects = content.get("projects") or []
        if 0 <= source_index < len(projects) and isinstance(projects[source_index], dict):
            projects[source_index]["description"] = generated_text.strip()
            content["projects"] = projects
            update_resume(resume, content=content)


def _normalize_priority(value: str) -> str:
    level = str(value or "").strip().lower()
    if level in {"high", "critical"}:
        return "high"
    if level in {"medium", "moderate"}:
        return "medium"
    return "low"


def _recommendation_sort_key(item: Dict[str, Any]):
    rank = {"high": 0, "medium": 1, "low": 2}
    return (rank.get(item.get("priority") or "low", 2), str(item.get("title") or ""))


def _build_resume_intelligence(
    content: Dict[str, Any],
    ats: Dict[str, Any],
    match_result: Optional[Dict[str, Any]],
    job_description: str,
) -> Dict[str, Any]:
    personal = content.get("personal_information") or {}
    experience = content.get("experience") or []
    projects = content.get("projects") or []
    skills = content.get("skills") or []

    recommendations: List[Dict[str, Any]] = []

    for rec in ats.get("recommendations") or []:
        if not isinstance(rec, dict):
            continue
        recommendations.append({
            "priority": _normalize_priority(rec.get("priority") or "medium"),
            "type": str(rec.get("type") or "general"),
            "title": str(rec.get("title") or "Improve Resume").strip(),
            "detail": str(rec.get("message") or "").strip(),
            "action": "review_section",
        })

    total_bullets = 0
    quantified_bullets = 0
    for entry in experience:
        if not isinstance(entry, dict):
            continue
        bullets = entry.get("bullets") or []
        for bullet in bullets:
            bullet_text = str(bullet).strip()
            if not bullet_text:
                continue
            total_bullets += 1
            if any(char.isdigit() for char in bullet_text):
                quantified_bullets += 1

    if content.get("summary") and len(str(content.get("summary") or "").split()) < 25:
        recommendations.append({
            "priority": "high",
            "type": "content",
            "title": "Strengthen Professional Summary",
            "detail": "Your summary is short. Add clearer value proposition, domain focus, and impact.",
            "action": "ai_rewrite_summary",
        })

    if experience and total_bullets < max(3, len(experience) * 2):
        recommendations.append({
            "priority": "high",
            "type": "experience",
            "title": "Add More Experience Evidence",
            "detail": "Each role should show impact through multiple bullet points.",
            "action": "add_experience_bullets",
        })

    if total_bullets and quantified_bullets / total_bullets < 0.35:
        recommendations.append({
            "priority": "medium",
            "type": "experience",
            "title": "Increase Quantified Achievements",
            "detail": "Add metrics, scope, or outcomes to bullet points for stronger credibility.",
            "action": "quantify_bullets",
        })

    if not projects:
        recommendations.append({
            "priority": "medium",
            "type": "projects",
            "title": "Add A Project Section",
            "detail": "Projects help demonstrate practical delivery and tool usage.",
            "action": "add_projects",
        })

    if not personal.get("professional_title"):
        recommendations.append({
            "priority": "medium",
            "type": "branding",
            "title": "Set A Professional Title",
            "detail": "Add a clear headline to instantly communicate your role focus.",
            "action": "update_personal_header",
        })

    missing_candidates: List[str] = []
    semantic_score = None
    candidate_match_score = None

    if isinstance(match_result, dict):
        semantic_score = match_result.get("semantic_score")
        candidate_match_score = match_result.get("candidate_match_score")
        missing_candidates = [str(item).strip() for item in (match_result.get("missing_candidates") or []) if str(item).strip()]

        if missing_candidates:
            recommendations.append({
                "priority": "high",
                "type": "keywords",
                "title": "Cover Missing Job Keywords",
                "detail": "Add evidence for top missing skills from the job description.",
                "action": "apply_missing_skills",
            })

        if isinstance(semantic_score, (int, float)) and semantic_score < 70:
            recommendations.append({
                "priority": "high",
                "type": "targeting",
                "title": "Tailor Resume To Job",
                "detail": "Your semantic alignment with this JD is low. Tailor summary and experience to role needs.",
                "action": "ai_tailor_resume",
            })

    seen_titles = set()
    deduped_recommendations = []
    for rec in recommendations:
        title_key = str(rec.get("title") or "").casefold().strip()
        if not title_key or title_key in seen_titles:
            continue
        seen_titles.add(title_key)
        deduped_recommendations.append(rec)

    deduped_recommendations.sort(key=_recommendation_sort_key)

    ats_score = float(ats.get("ats_score") or 0.0)
    readiness_score = ats_score
    if isinstance(semantic_score, (int, float)):
        readiness_score = round((ats_score * 0.65) + (float(semantic_score) * 0.35), 2)

    diagnostics = {
        "section_gaps": [
            section
            for section, present in {
                "summary": bool(content.get("summary")),
                "experience": bool(experience),
                "education": bool(content.get("education") or []),
                "skills": bool(skills),
                "projects": bool(projects),
            }.items()
            if not present
        ],
        "bullet_count": total_bullets,
        "quantified_bullet_count": quantified_bullets,
        "missing_keywords": missing_candidates[:12],
        "analysis_mode": (match_result or {}).get("analysis_mode") if isinstance(match_result, dict) else None,
    }

    return {
        "readiness_score": readiness_score,
        "ats_score": ats_score,
        "semantic_score": semantic_score,
        "candidate_match_score": candidate_match_score,
        "high_priority_count": sum(1 for rec in deduped_recommendations if rec.get("priority") == "high"),
        "recommendations": deduped_recommendations[:10],
        "diagnostics": diagnostics,
        "job_description_used": bool(job_description),
    }


def api_list_resumes():
    user_id = _require_api_user_id()
    if user_id is None:
        return _api_error(401, "Authentication required.", code="auth_required")

    resumes = [_resume_payload(resume) for resume in list_resumes_for_user(user_id)]
    return jsonify({
        "success": True,
        "resumes": resumes,
        "allowed_templates": sorted(ALLOWED_TEMPLATES),
    })


def api_create_resume():
    user_id = _require_api_user_id()
    if user_id is None:
        return _api_error(401, "Authentication required.", code="auth_required")

    body = request.get_json(silent=True) or {}
    creation_method = str(body.get("creation_method") or "scratch").strip().lower()
    title = body.get("title") or ""
    target_role = body.get("target_role") or ""

    try:
        if creation_method == "analysis":
            content = build_content_from_analysis(session)
            resume = create_resume(
                user_id,
                title=title,
                target_role=target_role,
                from_content=content,
            )
        elif creation_method == "duplicate":
            source_resume_id = body.get("source_resume_id")
            if not isinstance(source_resume_id, int):
                return _api_error(422, "source_resume_id is required for duplication.", code="invalid_payload")
            source_resume = get_owned_resume(source_resume_id, user_id)
            if source_resume is None:
                return _api_error(404, "Source resume not found.", code="not_found")
            resume = duplicate_resume(source_resume)
        elif creation_method == "ai_assisted":
            seed_content = body.get("content") if isinstance(body.get("content"), dict) else {}
            resume = create_resume(
                user_id,
                title=title or "AI Draft Resume",
                target_role=target_role,
                from_content=seed_content,
            )
        else:
            resume = create_resume(
                user_id,
                title=title,
                target_role=target_role,
                from_content=body.get("content") if isinstance(body.get("content"), dict) else None,
            )

        template = body.get("template")
        if template:
            if template not in ALLOWED_TEMPLATES:
                return _api_error(422, "Invalid template.", code="invalid_template")
            update_resume(resume, template=template)

        return jsonify({
            "success": True,
            "resume": _resume_payload(resume),
        }), 201
    except Exception:
        return _api_error(500, "Unable to create resume.", code="create_failed")


def api_get_resume(resume_id: int):
    user_id = _require_api_user_id()
    if user_id is None:
        return _api_error(401, "Authentication required.", code="auth_required")

    resume = get_owned_resume(resume_id, user_id)
    if resume is None:
        return _api_error(404, "Resume not found.", code="not_found")

    return jsonify({"success": True, "resume": _resume_payload(resume)})


def api_update_resume(resume_id: int):
    user_id = _require_api_user_id()
    if user_id is None:
        return _api_error(401, "Authentication required.", code="auth_required")

    resume = get_owned_resume(resume_id, user_id)
    if resume is None:
        return _api_error(403, "You cannot modify this resume.", code="ownership_violation")

    body = request.get_json(silent=True) or {}
    template = body.get("template")
    content = body.get("content")

    if template is not None and template not in ALLOWED_TEMPLATES:
        return _api_error(422, "Invalid template.", code="invalid_template")

    if content is not None and not isinstance(content, dict):
        return _api_error(422, "Invalid content payload.", code="invalid_payload")

    has_any_change = any(
        key in body for key in ("title", "target_role", "template", "content")
    )
    if not has_any_change:
        return _api_error(400, "No changes provided.", code="invalid_payload")

    try:
        update_resume(
            resume,
            title=body.get("title") if "title" in body else None,
            target_role=body.get("target_role") if "target_role" in body else None,
            template=template,
            content=content if isinstance(content, dict) else None,
        )

        if not bool(body.get("autosave")):
            create_resume_version(
                resume,
                action="manual_save",
                note="Manual save from builder workspace",
            )

        return jsonify({
            "success": True,
            "resume": _resume_payload(resume),
            "autosave": bool(body.get("autosave")),
        })
    except Exception:
        return _api_error(500, "Unable to save resume.", code="save_failed")


def api_delete_resume(resume_id: int):
    user_id = _require_api_user_id()
    if user_id is None:
        return _api_error(401, "Authentication required.", code="auth_required")

    resume = get_owned_resume(resume_id, user_id)
    if resume is None:
        return _api_error(403, "You cannot delete this resume.", code="ownership_violation")

    try:
        delete_resume(resume)
        return jsonify({"success": True})
    except Exception:
        return _api_error(500, "Unable to delete resume.", code="delete_failed")


def api_duplicate_resume(resume_id: int):
    user_id = _require_api_user_id()
    if user_id is None:
        return _api_error(401, "Authentication required.", code="auth_required")

    resume = get_owned_resume(resume_id, user_id)
    if resume is None:
        return _api_error(403, "You cannot duplicate this resume.", code="ownership_violation")

    try:
        clone = duplicate_resume(resume)
        return jsonify({"success": True, "resume": _resume_payload(clone)}), 201
    except Exception:
        return _api_error(500, "Unable to duplicate resume.", code="duplicate_failed")


def api_analyze_resume(resume_id: int):
    user_id = _require_api_user_id()
    if user_id is None:
        return _api_error(401, "Authentication required.", code="auth_required")

    resume = get_owned_resume(resume_id, user_id)
    if resume is None:
        return _api_error(403, "You cannot analyze this resume.", code="ownership_violation")

    body = request.get_json(silent=True) or {}
    job_description = str(body.get("job_description") or "").strip()

    content = sanitize_content(resume.get_content())
    text = _content_to_text(content)
    personal = content.get("personal_information") or {}

    if not text:
        return _api_error(422, "Resume content is empty.", code="empty_resume")

    try:
        ats = analyze_ats(
            resume_text=text,
            name=personal.get("full_name") or None,
            email=personal.get("email") or None,
            phone=personal.get("phone") or None,
            job_description=job_description or None,
        )

        match_result = None
        suggested_skills: List[Dict[str, Any]] = []
        if job_description:
            match_result = match_resume_to_job(text, job_description)
            existing = {str(skill).casefold() for skill in content.get("skills") or []}
            missing_candidates = match_result.get("missing_candidates") or []
            for candidate in missing_candidates:
                candidate_name = str(candidate).strip()
                if candidate_name and candidate_name.casefold() not in existing:
                    suggested_skills.append({
                        "name": candidate_name,
                        "suggested": True,
                        "reason": "Mentioned in job description but not evidenced in your resume yet.",
                    })

        intelligence = _build_resume_intelligence(content, ats, match_result, job_description)

        return jsonify({
            "success": True,
            "ats": ats,
            "match": match_result,
            "skill_suggestions": suggested_skills,
            "intelligence": intelligence,
            "recommendations": intelligence.get("recommendations") or [],
        })
    except ValueError as exc:
        return _api_error(422, str(exc), code="invalid_analysis_input")
    except Exception:
        return _api_error(500, "Unable to analyze resume.", code="analysis_failed")


def api_generate_resume_content(resume_id: int):
    user_id = _require_api_user_id()
    if user_id is None:
        return _api_error(401, "Authentication required.", code="auth_required")

    resume = get_owned_resume(resume_id, user_id)
    if resume is None:
        return _api_error(403, "You cannot generate content for this resume.", code="ownership_violation")

    body = request.get_json(silent=True) or {}
    operation = str(body.get("operation") or "").strip().lower()
    job_description = str(body.get("job_description") or "").strip()
    target_role = str(body.get("target_role") or resume.target_role or "").strip()

    content = sanitize_content(resume.get_content())
    resume_text = _content_to_text(content)
    source_text = _get_resume_source_for_generation(content, body)
    personal = content.get("personal_information") or {}

    resume_data = {
        "name": personal.get("full_name") or session.get("name") or "",
        "skills": content.get("skills") or [],
        "profession_field": personal.get("professional_title") or session.get("profession_field") or "Professional",
        "sections_found": content.get("section_order") or [],
        "word_count": len(resume_text.split()),
        "text": resume_text,
    }

    instruction_overrides = {
        "rewrite_summary": "Rewrite the summary to be more polished and professional.",
        "professional_summary": "Make the summary more professional and impactful.",
        "concise_summary": "Make the summary concise while preserving candidate facts.",
        "tailor_summary": "Tailor this summary for the job description using only candidate evidence.",
        "improve_summary_keywords": "Improve ATS keyword alignment in this summary using only true candidate facts.",
        "improve_experience": "Improve the impact and clarity of these bullets without inventing metrics.",
        "project_improve": "Improve this project description professionally without adding unsupported claims.",
        "rewrite_resume_professional": "Rewrite the full resume in a professional tone and preserve all facts.",
        "rewrite_resume_concise": "Rewrite the full resume to be concise and ATS-friendly while preserving facts.",
        "tailor_resume": "Tailor this resume to the job description without adding unsupported skills or achievements.",
    }

    task_map = {
        "generate_summary": TASK_SUMMARY,
        "rewrite_summary": TASK_SUMMARY,
        "professional_summary": TASK_SUMMARY,
        "concise_summary": TASK_SUMMARY,
        "tailor_summary": TASK_SUMMARY,
        "improve_summary_keywords": TASK_SUMMARY,
        "generate_experience_bullets": TASK_BULLETS,
        "improve_experience": TASK_EXPERIENCE,
        "project_improve": TASK_EXPERIENCE,
        "rewrite_resume_professional": TASK_RESUME_REWRITE,
        "rewrite_resume_concise": TASK_RESUME_REWRITE,
        "tailor_resume": TASK_RESUME_REWRITE,
    }

    task = task_map.get(operation)
    if task is None:
        return _api_error(422, "Unsupported AI operation.", code="invalid_operation")

    if operation in {"tailor_summary", "tailor_resume"} and not job_description:
        return _api_error(422, "Job description is required for tailoring.", code="missing_job_description")

    if task in {TASK_BULLETS, TASK_EXPERIENCE} and not source_text:
        return _api_error(422, "Source text is required for this operation.", code="missing_source_text")

    instructions = str(body.get("instructions") or "").strip()
    if not instructions and operation in instruction_overrides:
        instructions = instruction_overrides[operation]

    try:
        recommendations: List[Dict[str, Any]] = []
        if job_description:
            try:
                ats_preview = analyze_ats(
                    resume_text=resume_text,
                    name=personal.get("full_name") or None,
                    email=personal.get("email") or None,
                    phone=personal.get("phone") or None,
                    job_description=job_description or None,
                )
                match_preview = match_resume_to_job(resume_text, job_description)
                intelligence_preview = _build_resume_intelligence(content, ats_preview, match_preview, job_description)
                recommendations = intelligence_preview.get("recommendations") or []
            except Exception:
                recommendations = []

        apply_mode = body.get("apply") is True
        generated_content = str(body.get("generated_content") or "").strip()

        if apply_mode and generated_content:
            generated = generated_content
        else:
            generated = local_generate(
                task=task,
                resume_data=resume_data,
                source_text=source_text or None,
                job_description=job_description or None,
                target_role=target_role or None,
                instructions=instructions or None,
                recommendations=recommendations,
            )

        claim_validation = validate_generated_claims(resume_text, generated)

        if apply_mode:
            _apply_generated_content(resume, body, generated)
            create_resume_version(
                resume,
                action="ai_apply",
                note=f"AI apply operation: {operation}",
            )

        return jsonify({
            "success": True,
            "operation": operation,
            "content": generated,
            "warnings": claim_validation.get("warnings") or [],
            "unsupported_numeric_claims": claim_validation.get("unsupported_numeric_claims") or [],
            "recommendations_used": len(recommendations),
            "applied": bool(body.get("apply")),
            "resume": _resume_payload(resume),
        })
    except Exception:
        return _api_error(500, "Unable to generate content.", code="generation_failed")


def api_list_resume_versions(resume_id: int):
    user_id = _require_api_user_id()
    if user_id is None:
        return _api_error(401, "Authentication required.", code="auth_required")

    resume = get_owned_resume(resume_id, user_id)
    if resume is None:
        return _api_error(403, "You cannot access versions for this resume.", code="ownership_violation")

    versions = [
        {
            "id": version.id,
            "action": version.action,
            "note": version.note,
            "title": version.title,
            "template": version.template,
            "created_at": version.created_at.isoformat() if version.created_at else None,
        }
        for version in list_versions_for_resume(resume.id, user_id)
    ]

    return jsonify({"success": True, "versions": versions})


def api_create_resume_version(resume_id: int):
    user_id = _require_api_user_id()
    if user_id is None:
        return _api_error(401, "Authentication required.", code="auth_required")

    resume = get_owned_resume(resume_id, user_id)
    if resume is None:
        return _api_error(403, "You cannot version this resume.", code="ownership_violation")

    body = request.get_json(silent=True) or {}
    note = str(body.get("note") or "").strip()

    version = create_resume_version(
        resume,
        action="manual_version",
        note=note or "Manual version snapshot",
    )
    return jsonify({
        "success": True,
        "version": {
            "id": version.id,
            "action": version.action,
            "note": version.note,
            "created_at": version.created_at.isoformat() if version.created_at else None,
        },
    }), 201


def api_restore_resume_version(resume_id: int, version_id: int):
    user_id = _require_api_user_id()
    if user_id is None:
        return _api_error(401, "Authentication required.", code="auth_required")

    resume = get_owned_resume(resume_id, user_id)
    if resume is None:
        return _api_error(403, "You cannot restore this resume.", code="ownership_violation")

    version = get_owned_resume_version(version_id, user_id)
    if version is None or version.resume_id != resume.id:
        return _api_error(404, "Resume version not found.", code="version_not_found")

    restore_resume_from_version(resume, version)
    create_resume_version(
        resume,
        action="restore",
        note=f"Restored from version #{version.id}",
    )

    return jsonify({"success": True, "resume": _resume_payload(resume)})


def api_export_resume_pdf(resume_id: int):
    user_id = _require_api_user_id()
    if user_id is None:
        return _api_error(401, "Authentication required.", code="auth_required")

    resume = get_owned_resume(resume_id, user_id)
    if resume is None:
        return _api_error(403, "You cannot export this resume.", code="ownership_violation")

    try:
        payload = _export_pdf_bytes(resume)
        filename = f"{_filename_safe(resume.title)}.pdf"
        return send_file(
            io.BytesIO(payload),
            mimetype="application/pdf",
            as_attachment=True,
            download_name=filename,
        )
    except Exception:
        return _api_error(500, "Unable to export PDF.", code="pdf_export_failed")


def api_export_resume_docx(resume_id: int):
    user_id = _require_api_user_id()
    if user_id is None:
        return _api_error(401, "Authentication required.", code="auth_required")

    resume = get_owned_resume(resume_id, user_id)
    if resume is None:
        return _api_error(403, "You cannot export this resume.", code="ownership_violation")

    try:
        payload = _export_docx_bytes(resume)
        filename = f"{_filename_safe(resume.title)}.docx"
        return send_file(
            io.BytesIO(payload),
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            as_attachment=True,
            download_name=filename,
        )
    except Exception:
        return _api_error(500, "Unable to export DOCX.", code="docx_export_failed")
