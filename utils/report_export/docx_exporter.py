from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

from .report_data import ReportData


class DOCXReportExporter:
    """
    Generates a real DOCX report.

    Supported templates:
        professional
        modern
        minimal
        executive
        ai_insights
        career_analytics
    """

    SUPPORTED_TEMPLATES = {
        "professional",
        "modern",
        "minimal",
        "executive",
        "ai_insights",
        "career_analytics",
    }

    def export(
        self,
        report: ReportData,
        *,
        template: str = "professional",
    ) -> BytesIO:

        template = template.lower().strip()

        if template not in self.SUPPORTED_TEMPLATES:
            raise ValueError(
                f"Unsupported DOCX template: {template}"
            )

        document = Document()

        self._configure_page(document)

        accent = self._accent(template)

        self._add_header(
            document,
            report,
            accent,
        )

        self._add_score_table(
            document,
            report,
            accent,
        )

        self._add_section_heading(
            document,
            "ATS Breakdown",
            accent,
        )

        self._add_ats_table(
            document,
            report,
            accent,
        )

        self._add_section_heading(
            document,
            "Resume Metrics",
            accent,
        )

        self._add_metrics_table(
            document,
            report,
            accent,
        )

        self._add_section_heading(
            document,
            "Skills Analysis",
            accent,
        )

        self._add_skills(
            document,
            report,
        )

        self._add_section_heading(
            document,
            "Format Checker",
            accent,
        )

        self._add_format_checks(
            document,
            report,
        )

        self._add_section_heading(
            document,
            "Profile Completeness",
            accent,
        )

        self._add_profile(
            document,
            report,
        )

        self._add_recommendations(
            document,
            report,
            accent,
        )

        output = BytesIO()

        document.save(output)

        output.seek(0)

        return output

    # =========================================================
    # PAGE
    # =========================================================

    @staticmethod
    def _configure_page(
        document: Document,
    ) -> None:

        section = document.sections[0]

        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # =========================================================
    # HEADER
    # =========================================================

    def _add_header(
        self,
        document: Document,
        report: ReportData,
        accent: RGBColor,
    ) -> None:

        title = document.add_paragraph()

        title.alignment = WD_ALIGN_PARAGRAPH.LEFT

        run = title.add_run(
            "Resume Analysis Report"
        )

        run.bold = True
        run.font.size = Pt(22)
        run.font.color.rgb = accent

        subtitle = document.add_paragraph()

        run = subtitle.add_run(
            f"{report.name}  •  "
            f"{report.profession_field}"
        )

        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(
            102,
            112,
            133,
        )

        document.add_paragraph()

    # =========================================================
    # SCORE TABLE
    # =========================================================

    def _add_score_table(
        self,
        document: Document,
        report: ReportData,
        accent: RGBColor,
    ) -> None:

        table = document.add_table(
            rows=2,
            cols=4,
        )

        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        values = [
            f"{report.ats_score}/100",
            f"{report.match_percentage}%",
            str(len(report.skills)),
            f"{report.profile_strength}%",
        ]

        labels = [
            "ATS Score",
            "Job Match",
            "Skills Detected",
            "Profile Strength",
        ]

        for index in range(4):

            value_cell = table.cell(
                0,
                index,
            )

            value_cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            paragraph = value_cell.paragraphs[0]

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run(
                values[index]
            )

            run.bold = True
            run.font.size = Pt(14)
            run.font.color.rgb = accent

            label_cell = table.cell(
                1,
                index,
            )

            paragraph = label_cell.paragraphs[0]

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = paragraph.add_run(
                labels[index]
            )

            run.font.size = Pt(8)

    # =========================================================
    # ATS
    # =========================================================

    def _add_ats_table(
        self,
        document: Document,
        report: ReportData,
        accent: RGBColor,
    ) -> None:

        table = document.add_table(
            rows=1,
            cols=2,
        )

        table.style = "Table Grid"

        headers = [
            "Area",
            "Score",
        ]

        for index, header in enumerate(headers):

            run = table.cell(
                0,
                index,
            ).paragraphs[0].add_run(
                header
            )

            run.bold = True
            run.font.color.rgb = accent

        values = [
            ("Structure", "structure"),
            ("Content", "content"),
            ("Keywords", "keywords"),
            ("Contact", "contact"),
            ("Readability", "readability"),
        ]

        for label, key in values:

            cells = table.add_row().cells

            cells[0].text = label

            cells[1].text = (
                f"{report.ats_breakdown.get(key, 0)}%"
            )

    # =========================================================
    # METRICS
    # =========================================================

    def _add_metrics_table(
        self,
        document: Document,
        report: ReportData,
        accent: RGBColor,
    ) -> None:

        table = document.add_table(
            rows=2,
            cols=4,
        )

        table.style = "Table Grid"

        values = [
            str(report.word_count),
            str(report.page_count),
            f"{report.reading_time} min",
            f"{report.format_score}%",
        ]

        labels = [
            "Word Count",
            "Page Estimate",
            "Reading Time",
            "Format Score",
        ]

        for index in range(4):

            value = table.cell(
                0,
                index,
            ).paragraphs[0]

            value.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            run = value.add_run(
                values[index]
            )

            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = accent

            label = table.cell(
                1,
                index,
            ).paragraphs[0]

            label.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            label.add_run(
                labels[index]
            )

    # =========================================================
    # SKILLS
    # =========================================================

    @staticmethod
    def _add_skills(
        document: Document,
        report: ReportData,
    ) -> None:

        document.add_paragraph(
            "Detected Skills: "
            + (
                ", ".join(report.skills)
                if report.skills
                else "None detected"
            )
        )

        document.add_paragraph(
            "Matched Skills: "
            + (
                ", ".join(report.matched_skills)
                if report.matched_skills
                else "None detected"
            )
        )

        document.add_paragraph(
            "Missing Skills: "
            + (
                ", ".join(report.missing_skills)
                if report.missing_skills
                else "None detected"
            )
        )

    # =========================================================
    # FORMAT
    # =========================================================

    @staticmethod
    def _add_format_checks(
        document: Document,
        report: ReportData,
    ) -> None:

        table = document.add_table(
            rows=1,
            cols=3,
        )

        table.style = "Table Grid"

        headers = [
            "Check",
            "Status",
            "Detail",
        ]

        for index, header in enumerate(headers):

            run = table.cell(
                0,
                index,
            ).paragraphs[0].add_run(
                header
            )

            run.bold = True

        for item in report.format_checks:

            row = table.add_row().cells

            row[0].text = str(
                getattr(
                    item,
                    "label",
                    "",
                )
            )

            row[1].text = (
                "Passed"
                if getattr(
                    item,
                    "passed",
                    False,
                )
                else "Needs attention"
            )

            row[2].text = str(
                getattr(
                    item,
                    "detail",
                    "",
                )
            )

    # =========================================================
    # PROFILE
    # =========================================================

    @staticmethod
    def _add_profile(
        document: Document,
        report: ReportData,
    ) -> None:

        for item in report.profile_checklist:

            label = getattr(
                item,
                "label",
                "",
            )

            passed = getattr(
                item,
                "passed",
                False,
            )

            document.add_paragraph(
                f"{'✓' if passed else '○'} {label}"
            )

    # =========================================================
    # RECOMMENDATIONS
    # =========================================================

    def _add_recommendations(
        self,
        document: Document,
        report: ReportData,
        accent: RGBColor,
    ) -> None:

        if not report.recommendations:
            return

        self._add_section_heading(
            document,
            "AI Recommendations",
            accent,
        )

        for recommendation in report.recommendations:

            if isinstance(
                recommendation,
                dict,
            ):
                text = (
                    recommendation.get("message")
                    or recommendation.get("text")
                    or recommendation.get("recommendation")
                    or recommendation.get("title")
                    or str(recommendation)
                )
            else:
                text = str(
                    recommendation
                )

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                text
            )

    # =========================================================
    # SECTION HEADING
    # =========================================================

    @staticmethod
    def _add_section_heading(
        document: Document,
        title: str,
        accent: RGBColor,
    ) -> None:

        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            title
        )

        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = accent

    # =========================================================
    # TEMPLATE ACCENTS
    # =========================================================

    @staticmethod
    def _accent(
        template: str,
    ) -> RGBColor:

        colors_map = {
            "professional": (29, 78, 216),
            "modern": (15, 118, 110),
            "minimal": (52, 64, 84),
            "executive": (124, 58, 237),
            "ai_insights": (3, 105, 161),
            "career_analytics": (180, 83, 9),
        }

        return RGBColor(
            *colors_map.get(
                template,
                colors_map["professional"],
            )
        )


docx_exporter = DOCXReportExporter()