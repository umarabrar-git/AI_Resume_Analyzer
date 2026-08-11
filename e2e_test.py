import os
from app import app
from docx import Document

UPLOADS_DIR = os.path.join(os.getcwd(), "uploads")
os.makedirs(UPLOADS_DIR, exist_ok=True)

resume_path = os.path.join(UPLOADS_DIR, "test_resume.docx")

# Create a sample DOCX resume
resume_doc = Document()
resume_doc.add_paragraph("John Doe")
resume_doc.add_paragraph("Experienced software engineer with expertise in Python, Flask, Docker, SQL, and Git.")
resume_doc.add_paragraph("Skills:\n- Python\n- Flask\n- SQL\n- Git\n- Docker")
resume_doc.save(resume_path)

job_description = (
    "We are looking for a Python Developer.\n"
    "Responsibilities:\n"
    "- Develop web applications.\n"
    "- Collaborate with the engineering team.\n"
    "Skills: Python, Flask, SQL, Git, Docker"
)

with app.test_client() as client:
    with open(resume_path, "rb") as resume_file:
        response = client.post(
            "/upload",
            data={
                "resume": (resume_file, "test_resume.docx"),
                "job_description": job_description,
            },
            content_type="multipart/form-data",
        )

print("status_code=", response.status_code)
print("response_length=", len(response.get_data(as_text=True)))
print("contains_matched=", "Matched Skills" in response.get_data(as_text=True) or "✅" in response.get_data(as_text=True))
print("response_preview=\n", response.get_data(as_text=True)[:1000])

# Cleanup test resume
try:
    os.remove(resume_path)
except OSError:
    pass
